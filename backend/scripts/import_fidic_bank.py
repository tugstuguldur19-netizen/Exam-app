#!/usr/bin/env python3
"""Converts the "FIDIC Contract Manager — Practice Test Bank" .docx (300
questions across the 2017 Red, Yellow and Silver Books, answer key in tables)
into prisma/content/fidic/c01.json … c21.json — one file per clause, all
three books together.

    pip install python-docx
    python3 scripts/import_fidic_bank.py path/to/FIDIC_300_Practice_Tests.docx
    npm run seed

Question and option text is copied verbatim. Fails loudly (and writes
nothing) if any question, option, answer or clause count doesn't check out.
"""
import json
import re
import sys
from pathlib import Path

import docx

OUT_DIR = Path(__file__).resolve().parent.parent / "prisma" / "content" / "fidic"
BOOKS = {"A": "Red Book 2017", "B": "Yellow Book 2017", "C": "Silver Book 2017"}

PART = re.compile(r"^Part ([ABC]) — ")
CLAUSE = re.compile(r"^Clause (\d{1,2}) — (.+?)\s+\((\d+) questions?\)$")
QUESTION = re.compile(r"^(\d{1,3})\.\s+(.+)$")
OPTION = re.compile(r"^([a-e])\)\s+(.+)$")


def fail(msg):
    sys.exit(f"import_fidic_bank: {msg}")


def main(path):
    doc = docx.Document(path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # --- questions (everything between the first "Part A" and "Answer Key")
    try:
        start = next(i for i, t in enumerate(paras) if PART.match(t))
        end = paras.index("Answer Key")
    except (StopIteration, ValueError):
        fail("couldn't find the 'Part A' and 'Answer Key' headings")

    questions = {}  # number -> dict
    clause_expected = []  # (book, clause, count, [numbers])
    book = clause = None
    current = None
    for t in paras[start:end]:
        if m := PART.match(t):
            book, clause, current = BOOKS[m.group(1)], None, None
        elif m := CLAUSE.match(t):
            clause = int(m.group(1))
            clause_expected.append((book, clause, int(m.group(3)), []))
            current = None
        elif m := QUESTION.match(t):
            n = int(m.group(1))
            if n in questions:
                fail(f"question {n} appears twice")
            if clause is None:
                fail(f"question {n} is outside a clause section")
            current = {"n": n, "book": book, "clause": clause, "q": m.group(2), "options": []}
            questions[n] = current
            clause_expected[-1][3].append(n)
        elif m := OPTION.match(t):
            if current is None:
                fail(f"option '{t[:40]}' before any question")
            current["options"].append({"label": m.group(1).upper(), "text": m.group(2).strip()})
        else:
            fail(f"unrecognised line: {t[:80]!r}")

    # --- answer key tables: rows of (Q, Ans, Sub-Clause) × 3
    key = {}
    for table in doc.tables:
        header = [c.text.strip() for c in table.rows[0].cells]
        if header[:3] != ["Q", "Ans", "Sub-Clause"]:
            continue
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            for i in range(0, len(cells) - 2, 3):
                n, ans, ref = cells[i : i + 3]
                if not n:
                    continue
                if int(n) in key:
                    fail(f"answer key lists question {n} twice")
                key[int(n)] = (ans.lower(), ref)

    # --- checks
    if len(questions) != 300 or sorted(questions) != list(range(1, 301)):
        fail(f"expected questions 1–300, found {len(questions)}")
    if sorted(key) != sorted(questions):
        missing = sorted(set(questions) - set(key))
        fail(f"answer key doesn't match questions (missing: {missing[:10]})")
    for book_name, cl, count, nums in clause_expected:
        if len(nums) != count:
            fail(f"{book_name} clause {cl}: heading says {count} questions, found {len(nums)}")

    lessons = {c: [] for c in range(1, 22)}
    for n in sorted(questions):
        q = questions[n]
        if len(q["options"]) != 4:
            fail(f"question {n} has {len(q['options'])} options, expected 4")
        ans, ref = key[n]
        labels = [o["label"].lower() for o in q["options"]]
        if ans not in labels:
            fail(f"question {n}: answer '{ans}' is not one of its options")
        item = {
            "q": q["q"],
            "choices": q["options"],
            "answer": labels.index(ans),
            "explanation": f"Эх сурвалж: FIDIC {q['book']}, Sub-Clause {ref}.",
        }
        if not lessons[q["clause"]]:
            item["trial"] = True
        lessons[q["clause"]].append(item)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for c, items in lessons.items():
        (OUT_DIR / f"c{c:02d}.json").write_text(json.dumps(items, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"c{c:02d}.json: {len(items)} questions")
    print(f"Total: {sum(len(v) for v in lessons.values())} questions. Run `npm run seed` to load them.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
